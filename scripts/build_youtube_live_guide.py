from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "output/guides/takecast-youtube-live-setup-guide.docx"

NAVY = "211850"
PURPLE = "30246B"
TEAL = "008F82"
BLUE = "EAF0FF"
MINT = "E9F7F4"
WARM = "FFF4E6"
INK = "111111"
MUTED = "595959"
WHITE = "FFFFFF"
LIGHT_RULE = "D9D9E6"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:{}".format(key)), str(value))


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url, color=PURPLE, underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def paragraph_format(paragraph, before=0, after=6, line=1.08):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(doc, text, size=10, color=INK, bold=False, before=0, after=6, align=None):
    p = doc.add_paragraph()
    paragraph_format(p, before, after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_heading(doc, text, level=1, before=10, after=5):
    p = doc.add_paragraph()
    paragraph_format(p, before, after, 1.0)
    r = p.add_run(text)
    r.font.name = "Aptos Display"
    r.font.size = Pt(16 if level == 1 else 12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)
    return p


def add_label(doc, text):
    p = doc.add_paragraph()
    paragraph_format(p, 0, 4, 1.0)
    r = p.add_run(text.upper())
    r.font.name = "Aptos"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    return p


def add_bullet(doc, text, level=0, after=3):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph_format(p, 0, after, 1.08)
    r = p.add_run(text)
    r.font.name = "Aptos"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(INK)
    return p


def add_numbered_step(doc, number, title, body):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.columns[0].width = Inches(0.36)
    table.columns[1].width = Inches(6.62)
    left, right = table.rows[0].cells
    left.width = Inches(0.36)
    right.width = Inches(6.62)
    shade(left, PURPLE)
    shade(right, "F7F7FB")
    set_cell_margins(left, top=72, start=60, bottom=72, end=60)
    set_cell_margins(right, top=72, start=110, bottom=72, end=110)
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_border(cell, top={"val": "single", "sz": "6", "color": LIGHT_RULE}, bottom={"val": "single", "sz": "6", "color": LIGHT_RULE}, left={"val": "single", "sz": "6", "color": LIGHT_RULE}, right={"val": "single", "sz": "6", "color": LIGHT_RULE})
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_format(p, 0, 0, 1.0)
    run = p.add_run(str(number))
    run.font.name = "Aptos Display"
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(WHITE)
    p = right.paragraphs[0]
    paragraph_format(p, 0, 2, 1.0)
    r = p.add_run(title)
    r.font.name = "Aptos Display"
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(INK)
    p = right.add_paragraph()
    paragraph_format(p, 0, 0, 1.06)
    r = p.add_run(body)
    r.font.name = "Aptos"
    r.font.size = Pt(9.15)
    r.font.color.rgb = RGBColor.from_string(INK)
    p = doc.add_paragraph()
    paragraph_format(p, 0, 4, 1.0)
    return table


def add_note(doc, lead, text, fill=WARM):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=115, start=140, bottom=115, end=140)
    cell_border(cell, top={"val": "single", "sz": "7", "color": LIGHT_RULE}, bottom={"val": "single", "sz": "7", "color": LIGHT_RULE}, left={"val": "single", "sz": "7", "color": LIGHT_RULE}, right={"val": "single", "sz": "7", "color": LIGHT_RULE})
    p = cell.paragraphs[0]
    paragraph_format(p, 0, 0, 1.07)
    r = p.add_run(lead + " ")
    r.font.name = "Aptos"
    r.font.bold = True
    r.font.size = Pt(9.2)
    r.font.color.rgb = RGBColor.from_string(INK)
    r = p.add_run(text)
    r.font.name = "Aptos"
    r.font.size = Pt(9.2)
    r.font.color.rgb = RGBColor.from_string(INK)
    p = doc.add_paragraph()
    paragraph_format(p, 0, 4, 1.0)


def add_check(doc, text, after=3):
    p = doc.add_paragraph()
    paragraph_format(p, 0, after, 1.08)
    r = p.add_run("☐  ")
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(PURPLE)
    r = p.add_run(text)
    r.font.name = "Aptos"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(INK)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_header_footer(section):
    for header in (section.header, section.even_page_header):
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph_format(p, 0, 0, 1.0)
        r = p.add_run("TAKECAST  /  YOUTUBE LIVE")
        r.font.name = "Aptos"
        r.font.size = Pt(7.8)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(PURPLE)

    for footer in (section.footer, section.even_page_footer):
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph_format(p, 0, 0, 1.0)
        r = p.add_run("TakeCast YouTube Live Setup Guide   •   ")
        r.font.name = "Aptos"
        r.font.size = Pt(7.8)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        p._p.append(field)


def configure_document(doc):
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.top_margin = Inches(0.60)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    add_header_footer(section)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
    for name in ("List Bullet", "List Bullet 2"):
        styles[name].font.name = "Aptos"
        styles[name].font.size = Pt(9.5)


def title_page(doc):
    p = doc.add_paragraph()
    paragraph_format(p, 42, 4, 1.0)
    r = p.add_run("BEGINNER GUIDE")
    r.font.name = "Aptos"
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    p = doc.add_paragraph()
    paragraph_format(p, 0, 8, 1.0)
    r = p.add_run("TakeCast YouTube Live Setup Guide")
    r.font.name = "Aptos Display"
    r.font.bold = True
    r.font.size = Pt(27)
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    paragraph_format(p, 0, 17, 1.14)
    r = p.add_run("Enable YouTube Live, connect TakeCast, and run a safe first broadcast.")
    r.font.name = "Aptos"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    add_note(doc, "Start early.", "Your first YouTube Live activation can take up to 24 hours. Complete pages 1 and 2 at least one day before a real event.")
    add_heading(doc, "Start Here", before=11, after=5)
    add_text(doc, "The YouTube app and TakeCast do different jobs. The YouTube app lets you enable Live and create a direct phone stream. TakeCast is an encoder: it records a high-quality 4K copy on your iPhone and sends a live HLS feed to YouTube.", size=10, after=7)

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    headers = ("Use", "For")
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        paragraph_format(p, 0, 0, 1.0)
        r = p.add_run(text)
        r.font.name = "Aptos"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("YouTube app or YouTube Studio", "Verify your channel and enable live streaming."),
        ("TakeCast", "Record in 4K and send your HLS broadcast to the YouTube stream you prepared."),
    ]
    for row_index, (left_text, right_text) in enumerate(rows, start=1):
        for col, text in enumerate((left_text, right_text)):
            cell = table.rows[row_index].cells[col]
            shade(cell, BLUE if row_index % 2 else "FFFFFF")
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            paragraph_format(p, 0, 0, 1.05)
            r = p.add_run(text)
            r.font.name = "Aptos"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string(INK)
            if col == 0:
                r.font.bold = True
    p = doc.add_paragraph()
    paragraph_format(p, 0, 3, 1.0)

    add_heading(doc, "What You Need", before=8, after=4)
    for item in [
        "A YouTube channel you can manage, with a verified phone number and no current live-streaming restriction.",
        "An iPhone with TakeCast installed, enough storage for your 4K recording, power for the event, and a reliable internet connection.",
        "A computer or tablet browser for YouTube Studio. It is the easiest place to create an HLS stream key and inspect the live preview.",
    ]:
        add_bullet(doc, item)
    add_note(doc, "A quick distinction.", "YouTube’s mobile app currently has its own eligibility rules, including a subscriber threshold. This guide uses the encoder route that TakeCast needs; check YouTube’s current help if an eligibility message appears.", fill=MINT)


def enable_live_page(doc):
    add_label(doc, "Part One")
    add_heading(doc, "Enable YouTube Live", before=2, after=5)
    add_text(doc, "Do this before configuring TakeCast. You can begin in the YouTube app or in YouTube Studio. If Live is new to the channel, YouTube may hold activation for up to 24 hours.", size=10, after=8)
    add_numbered_step(doc, 1, "Verify the channel and turn on Live", "Open the YouTube app, tap Create, then Live. Follow the verification prompts. Or visit studio.youtube.com, choose Create, then Go live. If YouTube says activation is pending, wait for confirmation before moving on.")
    add_numbered_step(doc, 2, "Create an encoder stream in YouTube Studio", "In Live Control Room, open Stream. For a first stream, enter a title, select audience details, and create the stream. Use Unlisted for your first test so only people with the link can watch.")
    add_numbered_step(doc, 3, "Create an HLS stream key", "Under Select stream key, choose Create new stream key and select HLS as the protocol. Turn Auto start and Auto stop on. YouTube will change the Stream URL so it starts with https, not rtmp.")
    add_note(doc, "Important.", "The normal YouTube encoder screen often shows a server URL and a stream key separately. TakeCast needs the complete HLS ingest URL, not an RTMP address and not only a bare stream key.")
    add_heading(doc, "Copy the Right Credential", before=9, after=4)
    add_text(doc, "The HLS ingest URL is a long https address that includes your stream key. It commonly starts with https://a.upload.youtube.com/http_upload_hls?cid= and continues with secret values.", size=9.5, after=5)
    add_note(doc, "Treat it like a password.", "Anyone with the ingest URL can potentially broadcast on your channel. Send it only by a private method such as AirDrop, a password manager, or a direct message. Never put it in a public post, group chat, or support log.", fill=WARM)


def takecast_page(doc):
    add_label(doc, "Part Two")
    add_heading(doc, "Connect TakeCast and Test", before=2, after=5)
    add_text(doc, "Once YouTube Live is enabled and your HLS stream key exists, TakeCast only needs a one-time connection. A short private test is the best way to protect a real event.", size=10, after=8)
    add_numbered_step(doc, 4, "Connect TakeCast to YouTube", "Open TakeCast, choose Set up, turn on Stream to YouTube, and paste the complete HLS ingest URL. The field is hidden like a password. If you are unsure about video compression, keep the default H.264 setting.")
    add_numbered_step(doc, 5, "Run an unlisted test", "In YouTube Studio, keep the stream Unlisted. In TakeCast, open the camera, frame the shot, then tap REC. Wait for the live indicator to become healthy and confirm the picture in Live Control Room.")
    add_numbered_step(doc, 6, "End the test cleanly", "Tap STOP in TakeCast. Check that the local 4K recording appears under Review recordings and that YouTube has ended or archived the test. If the ingest URL was shared or exposed, reset the stream key in YouTube Studio.")
    add_heading(doc, "A Good First Test", before=8, after=4)
    for item in [
        "Use the same network, camera position, and power arrangement planned for the event.",
        "Record for at least one minute. Watch both the live preview and the locally saved recording.",
        "Use a separate Unlisted stream key for tests when possible. Keep the event key private.",
        "Share the channel live link with guests, such as youtube.com/@yourchannel/live, rather than one single broadcast link if you may need to restart.",
    ]:
        add_bullet(doc, item)
    add_note(doc, "Record only is safe.", "If Stream to YouTube is off or no valid ingest URL is saved, TakeCast still records the local 4K copy. Nothing is sent to YouTube.", fill=MINT)


def checklist_page(doc):
    add_label(doc, "Print This Page")
    add_heading(doc, "Event Day Checklist", before=2, after=6)
    add_text(doc, "Use this short list before guests arrive. Keep YouTube Studio open on a second screen if you can.", size=10, after=7)
    add_heading(doc, "Before You Press Record", level=2, before=4, after=4)
    for item in [
        "YouTube Live is enabled and the stream is set to the intended visibility.",
        "The correct HLS ingest URL is saved in TakeCast and Stream to YouTube is on.",
        "Phone has power, storage, and a tested internet connection.",
        "Camera is framed, audio is checked, and the landscape view looks right.",
        "A one-minute Unlisted test has succeeded at this location or on this network.",
    ]:
        add_check(doc, item)
    add_heading(doc, "During the Event", level=2, before=7, after=4)
    for item in [
        "Press REC in TakeCast and confirm the live indicator is healthy.",
        "Confirm the picture in YouTube Live Control Room before relying on the broadcast.",
        "Keep the phone powered and avoid moving it away from the tested network.",
    ]:
        add_check(doc, item)
    add_heading(doc, "If Something Goes Wrong", level=2, before=7, after=4)
    for item in [
        "Live is unavailable: complete channel verification and wait for YouTube to finish the first activation.",
        "TakeCast says Record only: turn on Stream to YouTube and confirm a complete HLS ingest URL is saved.",
        "No YouTube preview: make sure the stream key is HLS, not RTMP; then paste a newly copied complete HLS ingest URL.",
        "A key was exposed: reset it in YouTube Studio, then update TakeCast before broadcasting again.",
    ]:
        add_bullet(doc, item, after=2)
    add_heading(doc, "Official References", level=2, before=7, after=3)
    refs = [
        ("YouTube Help  Create a live stream with an encoder", "https://support.google.com/youtube/answer/2907883"),
        ("YouTube Help  Set up an HLS stream", "https://support.google.com/youtube/answer/10349430"),
        ("YouTube Help  Create a live stream on mobile", "https://support.google.com/youtube/answer/9228390"),
        ("Supplementary computer walkthrough", "https://themarketingheaven.com/how-to-go-live-on-youtube-from-computer/"),
    ]
    for label, url in refs:
        p = doc.add_paragraph()
        paragraph_format(p, 0, 2, 1.0)
        add_hyperlink(p, label, url)


def main():
    doc = Document()
    configure_document(doc)
    title_page(doc)
    add_page_break(doc)
    enable_live_page(doc)
    add_page_break(doc)
    takecast_page(doc)
    add_page_break(doc)
    checklist_page(doc)
    doc.core_properties.title = "TakeCast YouTube Live Setup Guide"
    doc.core_properties.subject = "Beginner instructions for enabling YouTube Live and connecting TakeCast"
    doc.core_properties.author = "TakeCast"
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
